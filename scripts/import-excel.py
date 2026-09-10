from __future__ import annotations

import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'sources' / 'locations-2026.xlsx'
OUTPUT = ROOT / 'src' / 'data' / 'locations.json'
CURATED = ROOT / 'src' / 'data' / 'curatedLocations.json'
VERIFIED_MEDIA = ROOT / 'src' / 'data' / 'verifiedMedia.json'
MEDIA_CHECK_DATE = '2026-09-08'

CATEGORY_MAP = {
    'Di tích lịch sử / khảo cổ': 'tourism',
    'Văn hóa / du lịch / tâm linh': 'tourism',
    'Lưu trú du lịch': 'lodging',
    'Ẩm thực / nhà hàng / cafe': 'food',
    'Vui chơi / thể thao': 'entertainment',
    'Y tế / chăm sóc sức khỏe': 'health',
    'Ngân hàng / ATM': 'utility',
    'Mua sắm / chợ / siêu thị': 'utility',
    'Bưu chính / viễn thông / điện / nhiên liệu': 'utility',
    'Hành chính / công an / PCCC': 'administration',
    'Giáo dục': 'utility',
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def text(value):
    value = clean(value)
    return None if value is None else str(value)


def slugify(value: str) -> str:
    normalized = unicodedata.normalize('NFD', value)
    ascii_text = ''.join(ch for ch in normalized if unicodedata.category(ch) != 'Mn')
    ascii_text = ascii_text.replace('đ', 'd').replace('Đ', 'D')
    ascii_text = re.sub(r'[^a-zA-Z0-9]+', '-', ascii_text).strip('-').lower()
    return ascii_text or 'dia-diem'


def load_json(path: Path, fallback):
    if not path.exists():
        return fallback
    return json.loads(path.read_text(encoding='utf-8'))


def normalize_media_fields(record: dict, verified_media: dict) -> dict:
    """Only attach media that has been explicitly verified for this exact location.

    No category-wide or generic fallback image is ever generated. If a place has not
    been verified, its image/gallery stay empty by design.
    """
    media = verified_media.get(record.get('id'))
    if media and media.get('image'):
        record.update({
            'image': media.get('image'),
            'imageAlt': media.get('imageAlt') or f"Ảnh {record.get('name', 'địa điểm')}",
            'gallery': media.get('gallery') or [],
            'imageSourceUrl': media.get('imageSourceUrl'),
            'imageSourceName': media.get('imageSourceName'),
            'imageVerification': 'verified_public',
            'imageCheckedAt': media.get('imageCheckedAt') or MEDIA_CHECK_DATE,
        })
    else:
        record['image'] = None
        record['imageAlt'] = None
        record['gallery'] = []
        record['imageVerification'] = 'not_verified'
        record['imageCheckedAt'] = MEDIA_CHECK_DATE
        record.pop('imageSourceUrl', None)
        record.pop('imageSourceName', None)
    return record


def main():
    wb = load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb['Danh_sach_chi_tiet']
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    index = {name: i for i, name in enumerate(headers)}

    rows = []
    used_ids = set()
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(values):
            continue

        def get(name):
            return clean(values[index[name]])

        stt = int(get('STT'))
        name = text(get('Tên địa điểm')) or f'Địa điểm {stt}'
        group = text(get('Nhóm')) or 'Khác'
        category = CATEGORY_MAP.get(group, 'utility')
        base_id = f"excel-{stt:03d}-{slugify(name)}"
        item_id = base_id
        suffix = 2
        while item_id in used_ids:
            item_id = f'{base_id}-{suffix}'
            suffix += 1
        used_ids.add(item_id)

        lat_raw = get('Latitude')
        lng_raw = get('Longitude')
        lat = float(lat_raw) if isinstance(lat_raw, (int, float)) else None
        lng = float(lng_raw) if isinstance(lng_raw, (int, float)) else None
        notes = text(get('Ghi chú'))
        if notes:
            notes = notes.replace('Google Travel/Maps', 'nguồn du lịch công khai').replace('Google Travel', 'nguồn du lịch công khai').replace('Google Maps', 'bản đồ công khai')

        record = {
            'id': item_id,
            'stt': stt,
            'name': name,
            'category': category,
            'group': group,
            'subgroup': text(get('Phân nhóm')),
            'address': text(get('Địa chỉ')),
            'lat': lat,
            'lng': lng,
            'phone': text(get('Điện thoại')),
            'websiteEmail': text(get('Website/Email')),
            'hours': text(get('Giờ hoạt động')),
            'keywords': text(get('Từ khóa / giá trị')),
            'heritageStatus': text(get('Tình trạng pháp lý di tích')),
            'notes': notes,
            'description': notes or '',
            'image': None,
            'imageAlt': None,
            'gallery': [],
            'panoramas': [],
            'imageVerification': 'not_verified',
            'imageCheckedAt': MEDIA_CHECK_DATE,
        }
        rows.append(record)

    curated = load_json(CURATED, {'overrides': {}, 'extras': []})
    verified_media = load_json(VERIFIED_MEDIA, {})
    overrides = curated.get('overrides', {})
    for record in rows:
        if record['id'] in overrides:
            record.update(overrides[record['id']])
        normalize_media_fields(record, verified_media)

    existing = {record['id'] for record in rows}
    for original in curated.get('extras', []):
        if not original.get('id') or original['id'] in existing:
            continue
        record = dict(original)
        normalize_media_fields(record, verified_media)
        rows.append(record)
        existing.add(record['id'])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(rows, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

    stats = {
        'total': len(rows),
        'withCoordinates': sum(1 for r in rows if r.get('lat') is not None and r.get('lng') is not None),
        'administration': sum(1 for r in rows if r.get('category') == 'administration'),
        'mediaAudited': sum(1 for r in rows if r.get('imageCheckedAt')),
        'verifiedImages': sum(1 for r in rows if r.get('imageVerification') == 'verified_public' and r.get('image')),
        'leftBlankByPolicy': sum(1 for r in rows if not r.get('image')),
    }
    print(json.dumps(stats, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
